from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import json
import math
import re

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "workbook" / "Week4_Customer_Support_Evaluation_Kajal_completed.xlsx"
METRICS_JSON = ROOT / "metadata" / "case_study_metrics.json"
RESULTS = ROOT / "evidence" / "results"
SCREENSHOTS = ROOT / "evidence" / "screenshots"
DIAGRAMS = ROOT / "assets" / "diagrams"
FIGURES = ROOT / "assets" / "figures"
DOCS = ROOT / "docs"
BUILD = ROOT / ".case-study-work"

NAVY = "16324F"
TEAL = "109D8E"
GREEN = "2E6B1F"
GOLD = "F2A51A"
RED = "D9534F"
INK = "202824"
MID = "5B6770"
GRID = "D9E2E1"
PALE_BLUE = "EEF4F8"
PALE_TEAL = "EAF7F4"
PALE_GOLD = "FFF6E2"
PALE_RED = "FCECEC"
WHITE = "FFFFFF"

LABELS = ["order_status", "refund_request", "product_issue", "account_help", "other"]
DISPLAY = {
    "order_status": "Order status",
    "refund_request": "Refund request",
    "product_issue": "Product issue",
    "account_help": "Account help",
    "other": "Other",
}

REPO_URL = "https://github.com/kajalchourasia-cmd/customer-support-agent-evals"
GDOC_URL = "https://docs.google.com/document/d/1n_LMh7LhJMKXgKXJWg9dqDNyviztHkASilYU_eq0Izk/edit"
LANGSMITH_URL = (
    "https://smith.langchain.com/o/b995b149-5185-4abe-9dd5-28045024bbb5/"
    "datasets/d14c5f9c-800e-4c32-a962-c112a36d3916/compare?selectedSessions="
    "33cdf37f-9c5f-4845-bddd-dc8143c47f7f%2C7e5de78b-82d0-4359-af2e-ce473e2e584f"
    "&source=33cdf37f-9c5f-4845-bddd-dc8143c47f7f"
)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    names = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for path in names:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def hex_rgb(value: str) -> tuple[int, int, int]:
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def chart_canvas(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((45, 45, 1555, 855), radius=28, fill="#F9FBFC", outline="#DCE6EA", width=2)
    draw.text((90, 82), title, font=load_font(42, True), fill=hex_rgb(NAVY))
    draw.text((90, 140), subtitle, font=load_font(23), fill=hex_rgb(MID))
    draw.line((90, 190, 1510, 190), fill=hex_rgb(TEAL), width=5)
    return image, draw


def draw_failure_taxonomy(path: Path) -> None:
    image, draw = chart_canvas(
        "Human-reviewed failure taxonomy",
        "All 29 failures in the supplied 150-row workbook",
    )
    data = [
        ("Refund intent under-routed", 12, TEAL),
        ("Account journey misrouted", 9, NAVY),
        ("Shipping context boundary", 4, GOLD),
        ("Product context boundary", 4, GREEN),
    ]
    left, top, max_width = 520, 265, 840
    for idx, (label, value, color) in enumerate(data):
        y = top + idx * 125
        draw.text((90, y + 11), label, font=load_font(28, True), fill=hex_rgb(INK))
        draw.rounded_rectangle((left, y, left + max_width, y + 58), radius=15, fill="#E8EEF1")
        width = int(max_width * value / 12)
        draw.rounded_rectangle((left, y, left + width, y + 58), radius=15, fill=hex_rgb(color))
        draw.text((left + width + 20, y + 8), str(value), font=load_font(34, True), fill=hex_rgb(INK))
    draw.text((90, 775), "Refund intent accounts for 41.4% of all reviewed failures.", font=load_font(24, True), fill=hex_rgb(NAVY))
    image.save(path)


def draw_grouped_metrics(path: Path, title: str, subtitle: str, series: list[tuple[str, list[float], str]]) -> None:
    image, draw = chart_canvas(title, subtitle)
    left, right, top, bottom = 160, 1490, 260, 720
    draw.line((left, top, left, bottom), fill="#91A1AA", width=2)
    draw.line((left, bottom, right, bottom), fill="#91A1AA", width=2)
    for tick in range(0, 101, 20):
        y = bottom - int((bottom - top) * tick / 100)
        draw.line((left, y, right, y), fill="#E2E8EA", width=1)
        draw.text((85, y - 15), f"{tick}%", font=load_font(22), fill=hex_rgb(MID))
    group_w = (right - left) / len(LABELS)
    bar_w = 46
    for group_idx, label in enumerate(LABELS):
        center = left + group_w * (group_idx + 0.5)
        total = len(series) * bar_w + (len(series) - 1) * 14
        start = center - total / 2
        for series_idx, (_, values, color) in enumerate(series):
            value = values[group_idx] * 100
            x0 = int(start + series_idx * (bar_w + 14))
            y0 = bottom - int((bottom - top) * value / 100)
            draw.rounded_rectangle((x0, y0, x0 + bar_w, bottom), radius=8, fill=hex_rgb(color))
            draw.text((x0 - 2, y0 - 30), f"{value:.0f}", font=load_font(19, True), fill=hex_rgb(INK))
        wrapped = DISPLAY[label].replace(" ", "\n")
        draw.multiline_text((center - 72, bottom + 25), wrapped, font=load_font(22, True), fill=hex_rgb(INK), align="center")
    legend_x = 1050
    for idx, (name, _, color) in enumerate(series):
        x = legend_x + idx * 220
        draw.rounded_rectangle((x, 215, x + 34, 249), radius=6, fill=hex_rgb(color))
        draw.text((x + 46, 216), name, font=load_font(22, True), fill=hex_rgb(INK))
    image.save(path)


def draw_notebook_outcomes(path: Path) -> None:
    image, draw = chart_canvas(
        "Frozen notebook experiment",
        "Aggregate improvement with complete paired outcome accounting",
    )
    # Accuracy cards
    for x, name, score, color in [(110, "Baseline V1", 91, NAVY), (470, "Focused V2", 98, TEAL)]:
        draw.rounded_rectangle((x, 245, x + 300, 545), radius=24, fill="#FFFFFF", outline=hex_rgb(color), width=4)
        draw.text((x + 48, 285), name, font=load_font(30, True), fill=hex_rgb(color))
        draw.text((x + 78, 355), f"{score}%", font=load_font(70, True), fill=hex_rgb(INK))
        draw.text((x + 70, 460), f"{100-score} errors", font=load_font(26), fill=hex_rgb(MID))
    draw.line((410, 395, 470, 395), fill=hex_rgb(TEAL), width=8)
    draw.polygon([(470, 395), (442, 378), (442, 412)], fill=hex_rgb(TEAL))
    # Outcome accounting
    draw.rounded_rectangle((860, 245, 1490, 655), radius=24, fill="#FFF9EC", outline=hex_rgb(GOLD), width=4)
    draw.text((915, 285), "Paired outcomes", font=load_font(34, True), fill=hex_rgb(NAVY))
    outcomes = [("Stable passes", 90, TEAL), ("Wrong to right", 8, GREEN), ("Right to wrong", 1, RED), ("Still wrong", 1, GOLD)]
    for idx, (label, value, color) in enumerate(outcomes):
        y = 355 + idx * 66
        draw.ellipse((920, y, 950, y + 30), fill=hex_rgb(color))
        draw.text((975, y - 3), label, font=load_font(27), fill=hex_rgb(INK))
        draw.text((1405, y - 3), str(value), font=load_font(28, True), fill=hex_rgb(INK), anchor="ra")
    draw.rounded_rectangle((110, 690, 1490, 790), radius=18, fill="#FCECEC", outline=hex_rgb(RED), width=2)
    draw.text((250, 715), "Regression rate = 1 / 91 = 1.10%   |   Pass gate ≤ 1.00%   |   Decision: REVIEW", font=load_font(28, True), fill=hex_rgb(RED))
    image.save(path)


def draw_confusion_matrices(path: Path, matrices: list[tuple[str, list[list[int]]]]) -> None:
    image, draw = chart_canvas(
        "Confusion matrices",
        "Rows are reference labels; columns are predicted labels",
    )
    short = ["order", "refund", "product", "account", "other"]
    starts = [100, 855]
    cell = 92
    for panel, (name, matrix) in enumerate(matrices):
        x0 = starts[panel]
        y0 = 340
        draw.text((x0, 225), name, font=load_font(34, True), fill=hex_rgb(NAVY))
        for col, label in enumerate(short):
            draw.text((x0 + 165 + col * cell, 285), label, font=load_font(19, True), fill=hex_rgb(MID), anchor="mm")
        for row, label in enumerate(short):
            draw.text((x0 + 80, y0 + row * cell + cell / 2), label, font=load_font(19, True), fill=hex_rgb(MID), anchor="mm")
            row_max = max(matrix[row]) or 1
            for col, value in enumerate(matrix[row]):
                intensity = value / row_max
                fill = (222 - int(95 * intensity), 244 - int(45 * intensity), 240 - int(35 * intensity)) if row == col else (255, 247 - int(35 * intensity), 230 - int(15 * intensity))
                x = x0 + 120 + col * cell
                y = y0 + row * cell
                draw.rounded_rectangle((x, y, x + cell - 8, y + cell - 8), radius=10, fill=fill, outline="#D6E0E3")
                draw.text((x + (cell - 8) / 2, y + (cell - 8) / 2), str(value), font=load_font(28, True), fill=hex_rgb(INK), anchor="mm")
    image.save(path)


def draw_operational_tradeoffs(path: Path) -> None:
    image, draw = chart_canvas(
        "Operational trade-off",
        "Focused V2 improved quality while increasing latency, tokens, and cost",
    )
    rows = [
        ("Mean latency", "0.773 s", "0.855 s", "+10.6%", TEAL),
        ("P95 latency", "1.028 s", "1.510 s", "+46.8%", GOLD),
        ("Mean tokens", "241.52", "476.46", "+97.3%", NAVY),
        ("Total cost / 100", "$0.00469", "$0.00830", "+77.0%", GREEN),
    ]
    headers = ["Measure", "Baseline V1", "Focused V2", "Relative change"]
    xs = [120, 650, 940, 1245]
    for x, label in zip(xs, headers):
        draw.text((x, 245), label, font=load_font(25, True), fill=hex_rgb(NAVY))
    draw.line((100, 290, 1500, 290), fill=hex_rgb(TEAL), width=4)
    for idx, (label, v1, v2, delta, color) in enumerate(rows):
        y = 330 + idx * 115
        if idx % 2:
            draw.rounded_rectangle((95, y - 18, 1505, y + 74), radius=12, fill="#F1F5F7")
        draw.text((120, y), label, font=load_font(28, True), fill=hex_rgb(INK))
        draw.text((650, y), v1, font=load_font(29), fill=hex_rgb(MID))
        draw.text((940, y), v2, font=load_font(29, True), fill=hex_rgb(INK))
        draw.rounded_rectangle((1240, y - 4, 1450, y + 48), radius=14, fill=hex_rgb(color))
        draw.text((1345, y + 22), delta, font=load_font(26, True), fill="white", anchor="mm")
    draw.text((120, 795), "Allowed-category validity remained 100% in both runs.", font=load_font(25, True), fill=hex_rgb(NAVY))
    image.save(path)


def draw_official_comparison(path: Path) -> None:
    image, draw = chart_canvas(
        "Separate workbook-aligned controlled comparison",
        "A distinct development experiment; not a continuation of the supplied 121/150 baseline",
    )
    cards = [
        (100, "Controlled V1", "145 / 150", "96.67%", NAVY),
        (480, "Focused V2", "146 / 150", "97.33%", TEAL),
    ]
    for x, label, count, pct, color in cards:
        draw.rounded_rectangle((x, 260, x + 330, 570), radius=25, fill="white", outline=hex_rgb(color), width=4)
        draw.text((x + 42, 305), label, font=load_font(31, True), fill=hex_rgb(color))
        draw.text((x + 48, 390), count, font=load_font(58, True), fill=hex_rgb(INK))
        draw.text((x + 92, 490), pct, font=load_font(30), fill=hex_rgb(MID))
    draw.line((430, 410, 480, 410), fill=hex_rgb(TEAL), width=8)
    draw.polygon([(480, 410), (452, 393), (452, 427)], fill=hex_rgb(TEAL))
    draw.rounded_rectangle((900, 260, 1490, 650), radius=25, fill="#FFF9EC", outline=hex_rgb(GOLD), width=4)
    draw.text((950, 305), "Outcome accounting", font=load_font(32, True), fill=hex_rgb(NAVY))
    lines = [("Stable passes", "145"), ("Wins", "1"), ("Regressions", "0"), ("Remaining failures", "4")]
    for i, (label, value) in enumerate(lines):
        y = 380 + i * 62
        draw.text((955, y), label, font=load_font(27), fill=hex_rgb(INK))
        draw.text((1420, y), value, font=load_font(29, True), fill=hex_rgb(INK), anchor="ra")
    draw.rounded_rectangle((100, 700, 1490, 800), radius=18, fill="#FCECEC", outline=hex_rgb(RED), width=2)
    draw.text((215, 725), "Decision: REVIEW because other precision = 0.882, below the 0.90 class-level gate.", font=load_font(29, True), fill=hex_rgb(RED))
    image.save(path)


def get_metrics() -> dict:
    if not METRICS_JSON.exists():
        raise FileNotFoundError("Run .case-study-work/compute_report_metrics.py first")
    return json.loads(METRICS_JSON.read_text(encoding="utf-8"))


def workbook_baseline() -> tuple[pd.DataFrame, pd.DataFrame, list[list[int]]]:
    wb = load_workbook(WORKBOOK, data_only=True)
    ws = wb["Step 1 · Tag failures"]
    records = []
    for row in range(15, 165):
        records.append({
            "id": ws.cell(row, 1).value,
            "ticket_text": ws.cell(row, 2).value,
            "true_category": ws.cell(row, 3).value,
            "predicted_category": ws.cell(row, 4).value,
            "reasoning": ws.cell(row, 5).value,
            "result": ws.cell(row, 6).value,
            "annotation": ws.cell(row, 7).value,
        })
    rows = pd.DataFrame(records)
    metrics = pd.DataFrame([
        {"category": ws.cell(r,1).value, "precision": ws.cell(r,2).value, "recall": ws.cell(r,3).value, "f1": ws.cell(r,4).value, "support": ws.cell(r,5).value}
        for r in range(4, 10)
    ])
    matrix = pd.crosstab(
        pd.Categorical(rows.true_category, categories=LABELS),
        pd.Categorical(rows.predicted_category, categories=LABELS),
        dropna=False,
    ).astype(int).values.tolist()
    return rows, metrics, matrix


def ensure_assets(metrics: dict, wb_metrics: pd.DataFrame, wb_matrix: list[list[int]]) -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    draw_failure_taxonomy(FIGURES / "01_failure_taxonomy.png")
    draw_grouped_metrics(
        FIGURES / "02_workbook_baseline_class_metrics.png",
        "Supplied workbook baseline by class",
        "Precision, recall, and F1 from the 150-row human-audit workbook",
        [
            ("Precision", wb_metrics[wb_metrics.category.isin(LABELS)].precision.tolist(), NAVY),
            ("Recall", wb_metrics[wb_metrics.category.isin(LABELS)].recall.tolist(), TEAL),
            ("F1", wb_metrics[wb_metrics.category.isin(LABELS)].f1.tolist(), GOLD),
        ],
    )
    draw_notebook_outcomes(FIGURES / "03_notebook_outcomes.png")
    v1_f1 = [metrics["notebook_v1"]["classification_report"][x]["f1-score"] for x in LABELS]
    v2_f1 = [metrics["notebook_v2"]["classification_report"][x]["f1-score"] for x in LABELS]
    draw_grouped_metrics(
        FIGURES / "04_notebook_class_f1.png",
        "Frozen notebook class-level F1",
        "Aggregate accuracy improved, while the product/other boundary remained imperfect",
        [("Baseline V1", v1_f1, NAVY), ("Focused V2", v2_f1, TEAL)],
    )
    draw_confusion_matrices(
        FIGURES / "05_notebook_confusion_matrices.png",
        [
            ("Baseline V1 — 91/100", metrics["notebook_v1"]["confusion_matrix"]),
            ("Focused V2 — 98/100", metrics["notebook_v2"]["confusion_matrix"]),
        ],
    )
    draw_operational_tradeoffs(FIGURES / "06_operational_tradeoffs.png")
    draw_official_comparison(FIGURES / "07_official_controlled_comparison.png")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_border(cell, color=GRID, size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = tc_borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_hyperlink(paragraph, text: str, url: str, color="0563C1"):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def set_keep(paragraph, with_next=False, together=False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if with_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if together:
        p_pr.append(OxmlElement("w:keepLines"))


def add_body(doc: Document, text: str, *, bold_lead: str | None = None, italic=False, space_after=6):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if italic:
        for run in p.runs:
            run.italic = True
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(4)


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep(p, with_next=True, together=True)
    return p


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    for run in p.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MID)
        run.italic = True
    set_keep(p, together=True)


def add_figure(doc: Document, path: Path, caption: str, width=6.75) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    set_keep(p, with_next=True, together=True)
    add_caption(doc, caption)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, label in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = str(label)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=110, start=120, bottom=110, end=120)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.color.rgb = RGBColor.from_string(WHITE)
                run.font.bold = True
                run.font.size = Pt(font_size)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = WHITE if r_idx % 2 == 0 else PALE_BLUE
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]
            cell.text = str(value)
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 and len(str(value)) < 30 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string(INK)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, before, after in [("Title", 28, 0, 8), ("Subtitle", 14, 0, 16), ("Heading 1", 18, 18, 7), ("Heading 2", 13, 12, 5), ("Heading 3", 11, 9, 4)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("000000")
        if name.startswith("Heading") or name == "Title":
            style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(10)

    header = section.header.paragraphs[0]
    header.text = "CUSTOMER SUPPORT AGENT EVALUATION  |  EVIDENCE FIRST CASE STUDY"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(MID)

    footer = section.footer.paragraphs[0]
    footer.text = "Kajal Chourasia"
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID)
    page_p = footer._parent.add_paragraph() if False else footer
    tab = footer.add_run("\t")
    tab.font.size = Pt(8)
    add_page_number(footer)


def pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def build_docx(metrics: dict, wb_rows: pd.DataFrame, wb_metrics: pd.DataFrame) -> Path:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph("Customer Support Agent Evaluation Case Study", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("A sequential, evidence-first evaluation of a five-class routing agent", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author = doc.add_paragraph("Kajal Chourasia")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.runs[0].bold = True
    author.paragraph_format.space_after = Pt(18)

    add_heading(doc, "Decision in one sentence", 1)
    add_body(doc, "Retain the focused prompt as a development candidate, but do not release it on the current evidence: the frozen notebook improved from 91% to 98%, yet its 1.10% regression rate missed the preregistered 1% gate; the separate workbook-aligned candidate also remains Review because other precision is 0.882, below the 0.90 class-level bar.")

    add_heading(doc, "Evidence at a glance", 1)
    add_table(
        doc,
        ["Evidence stream", "Cases", "Observed result", "Decision supported"],
        [
            ["Supplied workbook audit", "150", "121 correct; 29 reviewed failures", "Failure taxonomy and prompt hypothesis"],
            ["Frozen notebook comparison", "100", "91 → 98; 8 wins; 1 regression", "Focused V2 remains Review"],
            ["Workbook-aligned controlled run", "150", "145 → 146; 1 win; 0 regressions", "Separate candidate remains Review"],
        ],
        widths=[1.7, 0.6, 1.8, 2.5],
        font_size=8.4,
    )
    add_body(doc, "Evidence boundary. These three streams use different source artifacts or label contracts. Their scores are reported independently and are never combined into a single improvement claim.", bold_lead="Evidence boundary.")
    p = doc.add_paragraph()
    add_hyperlink(p, "GitHub repository", REPO_URL)
    p.add_run("  |  ")
    add_hyperlink(p, "Google Doc", GDOC_URL)
    p.add_run("  |  ")
    add_hyperlink(p, "LangSmith comparison", LANGSMITH_URL)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    add_heading(doc, "1. Product problem and evaluation objective", 1)
    add_body(doc, "The system routes a customer-support ticket to one of five operational queues. A wrong route can delay resolution, transfer the customer between teams, or send an agent into a workflow without the tools needed to solve the request. The evaluation therefore treats classification quality as a product decision, not as a leaderboard score.")
    add_heading(doc, "Routing contract", 2)
    add_table(
        doc,
        ["Category", "Operational meaning", "Boundary to protect"],
        [
            ["order_status", "Tracking, current location, delivery ETA, missing-delivery investigation", "Do not absorb general shipping questions or account-record problems"],
            ["refund_request", "Money back, credit, reversal, price adjustment, cancellation with returned funds", "Resolve explicit remedy intent according to the dataset's label policy"],
            ["product_issue", "Received item is damaged, wrong, defective, or not as described", "Separate post-purchase defects from pre-purchase questions"],
            ["account_help", "Login, password, saved address/payment method, checkout/account-access workflow", "Route by the blocked account journey, not nearby product/order nouns"],
            ["other", "General information, feedback, stock, policy, browsing, sales, or gift-card questions", "Avoid using other as a fallback for recognized operational failures"],
        ],
        widths=[1.05, 2.65, 3.0],
        font_size=8.2,
    )
    add_heading(doc, "Evaluation question", 2)
    add_body(doc, "Does a focused routing-policy change improve the highest-impact multi-intent failures while preserving previously correct behavior and keeping latency, token usage, and cost within acceptable bounds?")
    add_figure(doc, DIAGRAMS / "01_system_architecture.png", "Figure 1. System under evaluation and the separation between inference inputs and reference labels.")

    doc.add_page_break()
    add_heading(doc, "2. Evaluation design", 1)
    add_body(doc, "The workflow was designed around a repeatable product loop: define quality and risk, assemble and freeze evidence, run comparable experiments, analyze outcomes and traces, apply decision gates, then convert reviewed production failures into future regression cases.")
    add_figure(doc, DIAGRAMS / "02_evaluation_lifecycle.png", "Figure 2. Evaluation lifecycle from product intent to monitored production learning.")
    add_heading(doc, "Metrics and why they matter", 2)
    add_table(
        doc,
        ["Measure", "Question answered", "Decision use"],
        [
            ["Exact match", "Did the predicted route equal the reference route?", "Primary offline quality signal"],
            ["Precision / recall / F1", "Which queues are over- or under-routed?", "Find class-specific risk hidden by aggregate accuracy"],
            ["Allowed-category validity", "Did the model produce one valid route?", "Protect downstream automation and parsing"],
            ["Paired outcomes", "Which cases improved, regressed, stayed right, or stayed wrong?", "Prevent aggregate gains from hiding regressions"],
            ["Latency / tokens / cost", "What operating price accompanies the quality change?", "Assess customer experience and scale economics"],
            ["Trace review", "What input, reasoning, metadata, and component behavior caused a result?", "Diagnose mechanisms and verify auditability"],
        ],
        widths=[1.4, 2.5, 2.8],
        font_size=8.2,
    )
    add_heading(doc, "Release gates", 2)
    add_table(
        doc,
        ["Gate", "Pass condition", "Observed", "Status"],
        [
            ["Allowed output", "100% valid categories", "100% in both notebook runs", "PASS"],
            ["Net quality", "Candidate improves exact match", "91% → 98%", "PASS"],
            ["Regression budget", "Right-to-wrong rate ≤ 1% of V1-correct cases", "1 / 91 = 1.10%", "REVIEW"],
            ["Class floor", "Critical class precision ≥ 0.90", "other precision = 0.882 in workbook-aligned V2", "REVIEW"],
            ["External validity", "Independent holdout or production evidence", "Not available", "NOT ESTABLISHED"],
        ],
        widths=[1.25, 2.55, 1.75, 1.15],
        font_size=8.1,
    )

    add_heading(doc, "3. Evidence design and integrity controls", 1)
    add_figure(doc, DIAGRAMS / "03_evidence_integrity.png", "Figure 3. Three evidence streams and the conclusions each can support.")
    add_heading(doc, "Why the streams stay separate", 2)
    add_body(doc, "The supplied workbook audit evaluates historical predictions already present in the workbook. The frozen notebook comparison evaluates V1 and V2 on a separate 100-case generated dataset. The workbook-aligned controlled comparison runs two prompts on the fixed 150 workbook cases. Because the artifacts and label policies differ, 121/150 cannot be presented as the causal baseline for 146/150.")
    doc.add_page_break()
    add_heading(doc, "The key policy conflict", 2)
    add_table(
        doc,
        ["Ambiguous intent", "Workbook policy", "Notebook policy", "Implication"],
        [
            ["Missing order + explicit refund request", "refund_request", "order_status", "The same wording can have different correct labels"],
            ["Damaged item + explicit refund request", "refund_request for selected workbook cases", "product_issue", "A global refund-priority rule would damage notebook performance"],
            ["Promo code rejected at checkout", "account_help", "May look like other under generic policy", "Dataset-specific route ownership must be explicit"],
        ],
        widths=[1.55, 1.5, 1.5, 2.25],
        font_size=8.0,
    )
    add_body(doc, "Interpretation. This conflict is not a modeling footnote. It is a product-policy decision. A production system needs one adjudicated routing contract before scores from different benchmarks can be compared or one prompt can be optimized across them.", bold_lead="Interpretation.")
    add_heading(doc, "Integrity controls used", 2)
    add_bullets(doc, [
        "Stable case IDs and frozen V1/V2 exports for paired comparison.",
        "Reference labels used only by evaluators after inference, never passed to the classifier as answers.",
        "Prompt versions and SHA-256 hashes preserved with the result package.",
        "Every changed prediction reconciled into stable pass, win, regression, or remaining failure.",
        "Later benchmark-informed prompt versions, incomplete recovery runs, challenge-set experiments, and perfect-score claims excluded from the submission evidence.",
    ])

    add_heading(doc, "4. Step 1 — Audit the supplied 150-row workbook", 1)
    add_body(doc, "Every supplied prediction was reviewed against the ticket text and expected category. The workbook contained 121 passes and 29 failures, or 80.67% exact-match accuracy. The audit added a concrete human annotation to each failed row before grouping the failures.")
    wb_core = wb_metrics[wb_metrics.category.isin(LABELS)]
    add_table(
        doc,
        ["Reference class", "Precision", "Recall", "F1", "Support"],
        [[DISPLAY[r.category], f"{r.precision:.3f}", f"{r.recall:.3f}", f"{r.f1:.3f}", str(int(r.support))] for _, r in wb_core.iterrows()] +
        [["Macro average", "0.814", "0.825", "0.808", "150"]],
        widths=[2.0, 1.15, 1.15, 1.15, 1.0],
        font_size=8.8,
    )
    add_figure(doc, FIGURES / "02_workbook_baseline_class_metrics.png", "Figure 4. Supplied workbook baseline metrics by class.")
    add_heading(doc, "What the class metrics reveal", 2)
    add_bullets(doc, [
        "refund_request precision is 1.000 but recall is only 0.714: when the model predicts refund it is usually right, but it misses many true refund cases.",
        "order_status recall is high at 0.917 while precision is only 0.629, consistent with missing-order refund requests being absorbed into the status queue.",
        "account_help recall is 0.700, indicating that checkout, account-history, and promo-code failures are often routed by nearby nouns instead of the blocked workflow.",
        "Macro F1 is 0.808, confirming that the aggregate 80.67% accuracy is not hiding a single dominant high-performing class.",
    ])

    doc.add_page_break()
    add_heading(doc, "5. Steps 2 and 3 — Convert errors into a failure taxonomy", 1)
    add_body(doc, "The 29 annotations were clustered by decision mechanism rather than by surface wording. Every failed row was assigned to one of four mutually usable categories so the prompt hypothesis could target a policy boundary instead of memorizing examples.")
    add_figure(doc, FIGURES / "01_failure_taxonomy.png", "Figure 5. Distribution of the 29 human-reviewed workbook failures.")
    add_table(
        doc,
        ["Failure group", "Count", "Definition", "Product impact"],
        [
            ["Refund Intent Underrouted", "12", "Explicit refund, money-back, price-adjustment, or cancellation-plus-refund intent routed elsewhere", "Delays monetary resolution; adds transfers and escalation risk"],
            ["Account Journey Misrouted", "9", "Checkout, account visibility, or session failures routed by nearby product/order words", "Leaves customers blocked from payment or purchase records"],
            ["Shipping Context Boundary", "4", "Delivery incidents confused with profile changes or general shipping feedback", "Sends cases to teams with the wrong tools and SLA"],
            ["Product Context Boundary", "4", "Post-purchase claim failure or pre-purchase question assigned to the wrong side of the boundary", "Creates avoidable return work or leaves claims unresolved"],
        ],
        widths=[1.55, 0.55, 2.65, 2.1],
        font_size=7.9,
    )
    doc.add_page_break()
    add_heading(doc, "Representative annotated failures", 2)
    examples = [
        ["t012", "I never received my order and I want my money back.", "refund_request", "order_status", "Model followed missing-order context and ignored the requested monetary outcome."],
        ["t024", "The carrier marked it delivered but it shows the wrong delivery address.", "order_status", "account_help", "Model mapped 'address' to profile support although this is a delivery incident."],
        ["t103", "Site logs me out every time I try to pay.", "account_help", "product_issue", "Model treated the website malfunction as a product defect instead of a blocked checkout session."],
        ["t144", "Are the AirPods you sell authentic Apple products?", "other", "product_issue", "Model treated a pre-purchase authenticity question as a received-item defect."],
    ]
    add_table(doc, ["ID", "Ticket", "Reference", "Prediction", "Human diagnosis"], examples, widths=[0.5, 2.2, 0.95, 0.95, 2.15], font_size=7.3)

    add_heading(doc, "6. Step 4 — Form a focused prompt hypothesis", 1)
    add_body(doc, "Refund Intent Underrouted was selected because it was the largest group at 12 of 29 failures and because it directly affects money movement, trust, and escalation. The workbook prompt revision added one rule: explicit requests for a refund, money back, post-purchase price adjustment, or cancellation with returned funds should route to refund_request even when secondary context is present.")
    add_table(
        doc,
        ["Design element", "Baseline behavior", "Focused change", "Declared risk"],
        [
            ["Decision signal", "Keyword/topic proximity", "Use the customer's requested outcome", "Overweighting one outcome can damage other policy contracts"],
            ["Refund boundary", "Refund language could lose to shipping/account context", "Prioritize explicit monetary remedy under the workbook policy", "Status or account cases containing refund language may over-route"],
            ["Non-refund routes", "Broad definitions without precedence", "Keep status, replacement/repair, and account access when no monetary remedy is requested", "Ambiguous compound requests still require policy adjudication"],
            ["Prompt form", "Five label definitions", "One focused boundary rule; no case IDs or answer list", "General rule may still encode benchmark-specific policy"],
        ],
        widths=[1.2, 1.8, 2.15, 1.65],
        font_size=7.8,
    )
    add_heading(doc, "Pre-run hypothesis", 2)
    add_bullets(doc, [
        "Expected improvement: correct missing-order, price-adjustment, and subscription-cancellation cases labeled refund_request in the workbook.",
        "Regression risk: cases that mention refund language but are labeled by delivery or product context may move in the wrong direction.",
        "Evaluation safeguard: rerun on the same frozen IDs and inspect every changed prediction, not only the average score.",
    ])
    add_body(doc, "Important constraint. The notebook's 100-case label policy prioritizes underlying order or product context for several identical multi-intent patterns. The notebook V2 prompt therefore operationalizes its own frozen label contract. The two V2 results share a version name but are not one universal policy.", bold_lead="Important constraint.")

    doc.add_page_break()
    add_heading(doc, "7. Run the frozen notebook V1 and V2 comparison", 1)
    add_table(
        doc,
        ["Experiment property", "Value"],
        [
            ["Dataset", "100 frozen development examples; 20 references per class"],
            ["Case identity", "Stable IDs t000–t099 preserved across V1 and V2"],
            ["Model", "gpt-4o-mini as recorded in LangSmith run metadata"],
            ["Output schema", "Structured category, reasoning, and ticket text"],
            ["Evaluators", "Allowed-category validity and deterministic category exact match"],
            ["Tracing", "LangSmith root runs with dataset version, example ID, prompt version, model, runtime, latency, tokens, and cost"],
        ],
        widths=[1.7, 5.0],
        font_size=8.3,
    )
    add_figure(doc, FIGURES / "03_notebook_outcomes.png", "Figure 6. Frozen notebook aggregate results and paired outcome accounting.")
    add_heading(doc, "Aggregate result", 2)
    add_body(doc, "Baseline V1 scored 91/100 and focused V2 scored 98/100, a seven-percentage-point increase and a 77.8% relative reduction in observed errors. Both experiments returned an allowed category for all 100 examples.")
    add_table(
        doc,
        ["Outcome", "Cases", "Meaning"],
        [
            ["Stable pass", "90", "Correct in both versions"],
            ["Wrong to right", "8", "Candidate corrected a baseline failure"],
            ["Right to wrong", "1", "Candidate introduced a regression"],
            ["Still wrong", "1", "Neither version matched the reference"],
        ],
        widths=[1.55, 0.8, 4.35],
        font_size=8.8,
    )

    doc.add_page_break()
    add_heading(doc, "8. Inspect class-level movement", 1)
    v1_rep = metrics["notebook_v1"]["classification_report"]
    v2_rep = metrics["notebook_v2"]["classification_report"]
    class_rows = []
    for label in LABELS:
        f1a = v1_rep[label]["f1-score"]
        f1b = v2_rep[label]["f1-score"]
        note = {
            "order_status": "Four missing-order cases corrected",
            "refund_request": "Precision recovered from over-routing",
            "product_issue": "Improved, but two misleading-description cases remain wrong",
            "account_help": "Stable perfect score on this set",
            "other": "Recall stayed perfect; precision declined because two product cases moved here",
        }[label]
        class_rows.append([DISPLAY[label], f"{f1a:.3f}", f"{f1b:.3f}", f"{f1b-f1a:+.3f}", note])
    add_table(doc, ["Class", "V1 F1", "V2 F1", "Δ F1", "Interpretation"], class_rows, widths=[1.35, 0.75, 0.75, 0.7, 3.2], font_size=8.0)
    add_figure(doc, FIGURES / "04_notebook_class_f1.png", "Figure 7. Per-class F1 reveals the remaining product/other boundary risk.")
    add_figure(doc, FIGURES / "05_notebook_confusion_matrices.png", "Figure 8. Confusion matrices before and after the notebook prompt revision.")
    add_body(doc, "Reading the matrix. V1 over-routed four order-status and four product-issue references to refund_request, plus one product case to other. V2 repaired the refund over-routing, leaving two product_issue references predicted as other.", bold_lead="Reading the matrix.")

    doc.add_page_break()
    add_heading(doc, "9. Inspect every changed prediction", 1)
    changed = metrics["notebook_changed_cases"]
    changed_rows = []
    for item in changed:
        short_ticket = item["ticket_text"] if len(item["ticket_text"]) <= 83 else item["ticket_text"][:80] + "…"
        changed_rows.append([
            item["id"],
            short_ticket,
            DISPLAY[item["true_category"]],
            DISPLAY[item["predicted_category_v1"]],
            DISPLAY[item["predicted_category_v2"]],
            "WIN" if item["outcome"] == "wrong_to_right" else "REGRESSION",
        ])
    add_table(doc, ["ID", "Ticket", "Reference", "V1", "V2", "Outcome"], changed_rows, widths=[0.45, 2.55, 0.95, 0.85, 0.85, 0.95], font_size=6.9)
    add_heading(doc, "What changed mechanistically", 2)
    add_body(doc, "The eight wins are concentrated in two four-case families. Missing-order tickets labeled order_status stopped being captured by refund language. Damaged-product tickets labeled product_issue likewise stopped being captured by a requested refund. This concentration supports the intended precedence-rule mechanism on the notebook label policy.")
    add_body(doc, "The one regression is t052. The reference is product_issue because the purchased product failed an advertised wireless claim; V2 chose other because it interpreted the message as general product-information feedback. A paraphrase, t054, remained wrong for the same reason. The remaining risk is therefore a coherent product-description boundary, not random noise.")
    add_figure(doc, SCREENSHOTS / "02_eight_wins.png", "Figure 9. LangSmith comparison filtered to the eight wrong-to-right changes.")
    add_figure(doc, SCREENSHOTS / "03_one_regression.png", "Figure 10. LangSmith comparison filtered to the single right-to-wrong regression.")

    doc.add_page_break()
    add_heading(doc, "10. Connect aggregate scores to trace evidence", 1)
    add_body(doc, "Aggregate metrics show what changed; traces explain how the system behaved on a concrete request. The LangSmith runs preserve the model input, structured output, reasoning, evaluation scores, latency, token usage, cost, runtime, dataset version, example ID, and prompt version.")
    add_heading(doc, "Representative corrected trace", 2)
    add_figure(doc, SCREENSHOTS / "04_corrected_trace.png", "Figure 11. Corrected order-status trace with input, output, reasoning, and run metadata.", width=6.15)
    add_body(doc, "The corrected missing-order ticket now routes to order_status under the notebook policy. The trace makes the decision auditable and ties the row-level outcome to the aggregate eight-win count.")
    add_heading(doc, "Representative regression trace", 2)
    add_figure(doc, SCREENSHOTS / "05_regression_trace.png", "Figure 12. Regression trace for the misleading wireless-product-description ticket.", width=6.15)
    add_body(doc, "The regression trace preserves the counterexample needed for the next evaluation cycle. It should become a high-risk product-description slice in any future independently authored validation set.")

    doc.add_page_break()
    add_heading(doc, "11. Evaluate the operational trade-off", 1)
    op1 = metrics["notebook_operational_v1"]
    op2 = metrics["notebook_operational_v2"]
    rows = [
        ["Exact match", "91.0%", "98.0%", "+7.0 pp"],
        ["Allowed-category validity", "100.0%", "100.0%", "No change"],
        ["Mean latency", f"{op1['mean_latency_seconds']:.3f} s", f"{op2['mean_latency_seconds']:.3f} s", "+0.082 s"],
        ["P95 latency", f"{op1['p95_latency_seconds']:.3f} s", f"{op2['p95_latency_seconds']:.3f} s", "+0.481 s"],
        ["P99 latency", f"{op1['p99_latency_seconds']:.3f} s", f"{op2['p99_latency_seconds']:.3f} s", "+1.635 s"],
        ["Mean tokens", f"{op1['mean_tokens']:.2f}", f"{op2['mean_tokens']:.2f}", f"+{op2['mean_tokens']-op1['mean_tokens']:.2f}"],
        ["Total tokens", f"{int(op1['total_tokens']):,}", f"{int(op2['total_tokens']):,}", f"+{int(op2['total_tokens']-op1['total_tokens']):,}"],
        ["Total cost / 100", f"${op1['total_cost_usd']:.7f}", f"${op2['total_cost_usd']:.7f}", f"+${op2['total_cost_usd']-op1['total_cost_usd']:.7f}"],
    ]
    add_table(doc, ["Measure", "Baseline V1", "Focused V2", "Change"], rows, widths=[2.0, 1.55, 1.55, 1.55], font_size=8.6)
    add_figure(doc, FIGURES / "06_operational_tradeoffs.png", "Figure 13. Quality improved with a measurable latency, token, and cost increase.")
    add_body(doc, "Product interpretation. The absolute cost increase is tiny for 100 runs, but relative cost rises about 77% and mean tokens nearly double. At production scale, the decision should be based on avoided misroutes per incremental dollar and the customer impact of higher tail latency. The current data supports keeping V2 as a candidate, not deploying it without additional validation.", bold_lead="Product interpretation.")

    doc.add_page_break()
    add_heading(doc, "12. Apply the release decision", 1)
    add_figure(doc, DIAGRAMS / "04_release_decision.png", "Figure 14. The regression gate converts the 98% aggregate score into a Review decision.")
    add_table(
        doc,
        ["Decision criterion", "Evidence", "Assessment"],
        [
            ["Meaningful quality gain", "+7 percentage points; eight wins; error count 9 → 2", "Meets"],
            ["No invalid outputs", "100/100 allowed categories in both runs", "Meets"],
            ["Regression rate within budget", "1 regression / 91 previously correct = 1.10%; gate ≤ 1%", "Misses narrowly"],
            ["Remaining failure understood", "t052 regression and t054 remaining failure share one product-description boundary", "Understood; unresolved"],
            ["Operational cost acceptable", "+10.6% mean latency; +46.8% p95; +97.3% mean tokens; +77.0% total cost", "Requires product trade-off"],
            ["Independent generalization", "No organizer-hidden or independently authored holdout result", "Not established"],
        ],
        widths=[1.75, 3.65, 1.3],
        font_size=8.0,
    )
    add_heading(doc, "Decision — Review", 2)
    add_body(doc, "Retain V2 as a development candidate because it materially improves the frozen notebook set and its remaining errors are interpretable. Do not release solely on this evidence. The candidate misses the explicit regression gate, increases operating cost, and has not been evaluated on an independently authored holdout.")

    doc.add_page_break()
    add_heading(doc, "13. Separate workbook-aligned controlled comparison", 1)
    add_body(doc, "A second controlled experiment ran the original and focused workbook prompts on the same fixed 150 workbook cases. Controlled V1 scored 145/150 and focused V2 scored 146/150. This experiment is reported separately from the supplied workbook's historical 121/150 predictions because 25 supplied predictions changed when the original prompt was rerun under the current model/runtime.")
    add_figure(doc, FIGURES / "07_official_controlled_comparison.png", "Figure 15. Workbook-aligned controlled V1/V2 comparison and decision.")
    add_table(
        doc,
        ["Measure", "Controlled V1", "Focused V2", "Change"],
        [
            ["Exact match", "145 / 150 (96.67%)", "146 / 150 (97.33%)", "+1 case / +0.67 pp"],
            ["Macro F1", "0.966", "0.975", "+0.008"],
            ["Wins / regressions", "—", "1 / 0", "Positive paired result"],
            ["Mean latency", "0.740 s", "0.771 s", "+0.031 s"],
            ["Mean tokens", "231.49", "298.14", "+66.65"],
            ["Total cost", "$0.006843", "$0.008387", "+$0.001543"],
        ],
        widths=[1.65, 1.75, 1.75, 1.55],
        font_size=8.4,
    )
    add_heading(doc, "What actually improved", 2)
    add_body(doc, "The only wrong-to-right change was t025: a carrier delivered a parcel to the wrong address. Controlled V1 predicted product_issue; focused V2 predicted order_status, matching the reference. The intended refund cluster already had little headroom in controlled V1, so this experiment does not prove a large refund-policy effect.")
    doc.add_page_break()
    add_heading(doc, "Four remaining failures", 2)
    remaining = metrics["official_remaining_failures"]
    add_table(
        doc,
        ["ID", "Reference", "V2 prediction", "Failure mechanism"],
        [[x["id"], DISPLAY[x["true_category"]], DISPLAY[x["predicted_category_focused_v2"]], x["failure_group_from_supplied_workbook"].replace("_", " ")] for x in remaining],
        widths=[0.7, 1.35, 1.35, 3.3],
        font_size=8.6,
    )
    add_body(doc, "Decision. Focused V2 remains Review because other precision is 0.882, below the 0.90 class-level bar. The controlled comparison is useful reproducibility and robustness evidence, but its +1 result is modest and should not be described as a direct improvement from 121/150 to 146/150.", bold_lead="Decision.")

    add_heading(doc, "14. What the evaluation establishes — and what it does not", 1)
    add_table(
        doc,
        ["Claim", "Evidence status", "Reason"],
        [
            ["The supplied model has recurring policy-boundary failures", "Established", "All 29 supplied failures were manually annotated and grouped"],
            ["Notebook V2 improves the frozen 100-case development set", "Established", "Same IDs and references; 8 wins, 1 regression, 1 remaining failure"],
            ["Workbook-focused V2 modestly improves its controlled V1", "Established", "Same 150 cases; 1 win and 0 regressions"],
            ["121/150 causally becomes 146/150", "Not supported", "The 121 score is a different historical prediction artifact"],
            ["The prompt generalizes to unseen organizer or production data", "Not established", "No independent hidden or production set is available"],
            ["The optional LLM judge is calibrated and production-ready", "Not claimed", "The workbook contains scaffolding, but the judge extension was not completed as submission evidence"],
        ],
        widths=[2.45, 1.25, 3.0],
        font_size=7.9,
    )
    add_heading(doc, "Threats to validity", 2)
    add_table(
        doc,
        ["Threat", "Observed condition", "Mitigation / next test"],
        [
            ["Label-contract drift", "Workbook and notebook disagree on multi-intent precedence", "Adjudicate one production routing policy before cross-set comparison"],
            ["Development-set dependence", "Prompt hypotheses were derived from reviewed failures", "Freeze a new independently authored holdout before further tuning"],
            ["Seed/paraphrase dependence", "Notebook set contains related variants", "Report cluster-level confidence and evaluate semantically independent cases"],
            ["Model/runtime variability", "Historical workbook predictions differ from a fresh original-prompt run", "Pin model snapshot, parameters, package versions, and prompt hash"],
            ["Sparse tail evidence", "Only one observed notebook regression", "Use confidence intervals and add high-risk slices before release"],
            ["Offline-to-online gap", "No real routing outcomes or transfer costs", "Shadow deploy, sample traces, and measure resolution and transfer metrics"],
        ],
        widths=[1.5, 2.45, 2.75],
        font_size=7.8,
    )
    add_heading(doc, "Why the excluded evidence is excluded", 2)
    add_body(doc, "Later prompt versions were informed by observed benchmark behavior, some clean-room recovery runs were incomplete because LangSmith rejected trace ingestion, and challenge cases were designer-visible. Those artifacts can inform future work, but using them to claim blind generalization or a perfect final score would overstate the evidence. The submission retains only results with defensible provenance and complete paired accounting.")

    doc.add_page_break()
    add_heading(doc, "15. Recommended production evaluation plan", 1)
    add_numbered(doc, [
        "Publish one adjudicated routing policy with examples for every known cross-category boundary.",
        "Freeze the current candidate prompt, model snapshot, parameters, dependencies, and output schema before viewing any new labels.",
        "Commission an independently authored holdout with one row per semantic intent rather than paraphrase families, plus deliberate high-risk slices.",
        "Use two human reviewers for ambiguous cases and adjudicate disagreements before model scoring.",
        "Run the candidate once and retain the result whether it passes or fails. Report paired outcomes, per-class metrics, uncertainty, latency, tokens, and cost.",
        "Calibrate any LLM judge against human labels before using it as a release signal; report disagreement slices rather than judge accuracy alone.",
        "Shadow deploy before automatic routing. Sample traces, monitor route overrides and transfers, and maintain a human annotation queue.",
    ])
    add_heading(doc, "Suggested online scorecard", 2)
    add_table(
        doc,
        ["Signal", "Definition", "Suggested action threshold"],
        [
            ["Human override rate", "Share of routed tickets whose category is changed by an agent", "Investigate any class with sustained increase"],
            ["Transfer rate", "Tickets transferred after initial routing", "Prioritize high-volume or high-severity routes"],
            ["Time to correct queue", "Elapsed time until the ticket reaches the resolving workflow", "Alert on tail-latency and category regressions"],
            ["Resolution / reopen rate", "Ticket resolved without reopening", "Compare candidate against current routing policy"],
            ["Invalid or low-confidence output", "Schema failure or confidence below routing threshold", "Route to human triage; block automation"],
            ["Cost and latency", "Per-ticket token cost and p50/p95 response time", "Budget by traffic volume and SLA"],
        ],
        widths=[1.55, 2.75, 2.4],
        font_size=8.0,
    )

    doc.add_page_break()
    add_heading(doc, "16. Final conclusion", 1)
    add_body(doc, "The strongest result is not the 98% score by itself. It is the combination of a measurable seven-point gain, full paired outcome accounting, explicit recognition of one regression, class-level diagnosis of the remaining boundary, operational cost analysis, and a release decision that honors a preregistered gate.")
    add_body(doc, "The work also exposed a deeper product lesson: evaluation quality depends on policy quality. The same compound request is labeled differently across the workbook and notebook. Until the business decides which workflow owns that request, a model can be penalized for following one reasonable policy and rewarded for another. The next investment should therefore combine policy adjudication with an independent holdout, rather than further tuning on visible development cases.")
    add_body(doc, "Final recommendation. Keep focused V2 as a documented candidate. Do not claim hidden-test or production readiness. Resolve the routing contract, freeze the candidate, evaluate once on independent data, and require the regression, class-floor, and operating-cost gates to pass before release.", bold_lead="Final recommendation.")
    add_heading(doc, "Artifact and evidence index", 2)
    add_table(
        doc,
        ["Artifact", "Purpose", "Repository path"],
        [
            ["Completed workbook", "150-case human audit, taxonomy, labels, prompt hypothesis", "workbook/Week4_Customer_Support_Evaluation_Kajal_completed.xlsx"],
            ["Executed notebook", "Frozen V1/V2 workflow and saved outputs", "notebook/week4_customer_support_evals_executed.ipynb"],
            ["Paired notebook export", "Case-level wins, regressions, stable passes, failures", "evidence/results/baseline_vs_improved.csv"],
            ["Notebook validation summary", "Counts, IDs, hashes, and failure records", "evidence/results/validation_summary.json"],
            ["Workbook-aligned comparison", "150-case controlled V1/V2 outputs and metrics", "evidence/results/official150_case_comparison.csv"],
            ["LangSmith screenshots", "Aggregate, transitions, and trace-level evidence", "evidence/screenshots/"],
            ["Prompt versions", "Executed notebook and workbook-focused prompt policies", "prompts/"],
            ["Integrity metadata", "Package manifest and SHA-256 records", "metadata/"],
        ],
        widths=[1.55, 2.85, 2.3],
        font_size=7.4,
    )
    p = doc.add_paragraph()
    p.add_run("Live links: ").bold = True
    add_hyperlink(p, "Repository", REPO_URL)
    p.add_run("  |  ")
    add_hyperlink(p, "Google Doc", GDOC_URL)
    p.add_run("  |  ")
    add_hyperlink(p, "LangSmith comparison", LANGSMITH_URL)

    BUILD.mkdir(parents=True, exist_ok=True)
    out = BUILD / "Customer_Support_Evaluation_Case_Study_Detailed.docx"
    doc.save(out)
    return out


def build_markdown(metrics: dict) -> Path:
    path = DOCS / "Customer_Support_Evaluation_Case_Study.md"
    v1 = metrics["notebook_v1"]["classification_report"]
    v2 = metrics["notebook_v2"]["classification_report"]
    lines = [
        "# Customer Support Agent Evaluation Case Study",
        "",
        "**Author:** Kajal Chourasia",
        "**Decision:** **REVIEW**",
        "",
        "> Retain focused V2 as a development candidate. Do not release it on the current evidence: the frozen notebook improved from 91% to 98%, but the 1/91 = 1.10% regression rate narrowly missed the 1% gate.",
        "",
        "## 1. Product problem",
        "",
        "The agent routes each ticket to `order_status`, `refund_request`, `product_issue`, `account_help`, or `other`. A wrong route can delay resolution, create transfers, or trigger the wrong operating workflow. The evaluation therefore combines accuracy, class-level quality, regressions, trace evidence, latency, tokens, cost, and release gates.",
        "",
        "![System architecture](../assets/diagrams/01_system_architecture.png)",
        "",
        "## 2. Evaluation design",
        "",
        "The evaluation follows a frozen, paired comparison. Reference labels grade outputs after inference and are not given to the model as answers.",
        "",
        "![Evaluation lifecycle](../assets/diagrams/02_evaluation_lifecycle.png)",
        "",
        "| Measure | Purpose |",
        "|---|---|",
        "| Exact match | Primary offline routing-quality signal |",
        "| Per-class precision, recall, F1 | Expose over- and under-routing by queue |",
        "| Paired outcomes | Count wins, regressions, stable passes, and remaining failures |",
        "| Allowed-category validity | Protect downstream automation |",
        "| Latency, tokens, cost | Quantify the operating price of quality |",
        "| Trace review | Connect aggregate scores to concrete model behavior |",
        "",
        "## 3. Evidence boundaries",
        "",
        "![Evidence integrity](../assets/diagrams/03_evidence_integrity.png)",
        "",
        "| Evidence stream | Cases | Result | Supports |",
        "|---|---:|---:|---|",
        "| Supplied workbook audit | 150 | 121 correct; 29 reviewed failures | Failure taxonomy and prompt hypothesis |",
        "| Frozen notebook comparison | 100 | 91 → 98; 8 wins; 1 regression | Controlled development comparison |",
        "| Workbook-aligned controlled run | 150 | 145 → 146; 1 win; 0 regressions | Separate development evidence |",
        "",
        "The workbook and notebook disagree about several multi-intent labels. Their scores are never merged, and the controlled 145-to-146 result is not presented as the supplied workbook moving from 121 to 146.",
        "",
        "## 4. Human audit of the supplied workbook",
        "",
        "All 150 supplied predictions were reviewed. The workbook baseline is 121/150, or 80.67%, with macro F1 0.808.",
        "",
        "![Workbook class metrics](../assets/figures/02_workbook_baseline_class_metrics.png)",
        "",
        "## 5. Failure taxonomy",
        "",
        "![Failure taxonomy](../assets/figures/01_failure_taxonomy.png)",
        "",
        "| Failure group | Count | Product implication |",
        "|---|---:|---|",
        "| Refund Intent Underrouted | 12 | Monetary requests lose to surrounding context |",
        "| Account Journey Misrouted | 9 | Blocked checkout/account workflows route elsewhere |",
        "| Shipping Context Boundary | 4 | Delivery incidents and generic shipping feedback cross routes |",
        "| Product Context Boundary | 4 | Post-purchase claims and pre-purchase questions are confused |",
        "",
        "## 6. Focused prompt hypothesis",
        "",
        "The workbook prompt prioritizes an explicit requested monetary outcome under the workbook policy. The revision is a policy-level rule, not a list of case IDs or answers. Its declared regression risk is over-routing status, product, or account cases that happen to mention refund language.",
        "",
        "## 7. Frozen notebook result",
        "",
        "![Notebook outcomes](../assets/figures/03_notebook_outcomes.png)",
        "",
        "The frozen 100-case notebook moved from 91/100 to 98/100. Outcome accounting is 90 stable passes, 8 wins, 1 regression, and 1 remaining failure.",
        "",
        "![Notebook class F1](../assets/figures/04_notebook_class_f1.png)",
        "",
        "![Notebook confusion matrices](../assets/figures/05_notebook_confusion_matrices.png)",
        "",
        "The eight wins repaired two four-case families: missing-order references labeled `order_status` and damaged-product references labeled `product_issue`. The regression and remaining failure are two misleading wireless-description cases predicted as `other` instead of `product_issue`.",
        "",
        "## 8. Trace evidence",
        "",
        "![Eight wins](../evidence/screenshots/02_eight_wins.png)",
        "",
        "![One regression](../evidence/screenshots/03_one_regression.png)",
        "",
        "LangSmith ties the aggregate result to ticket input, structured output, reasoning, evaluation scores, latency, tokens, cost, runtime, dataset version, example ID, and prompt version.",
        "",
        "## 9. Operational trade-off",
        "",
        "![Operational trade-off](../assets/figures/06_operational_tradeoffs.png)",
        "",
        "| Measure | V1 | V2 | Change |",
        "|---|---:|---:|---:|",
        "| Mean latency | 0.773 s | 0.855 s | +10.6% |",
        "| P95 latency | 1.028 s | 1.510 s | +46.8% |",
        "| Mean tokens | 241.52 | 476.46 | +97.3% |",
        "| Total cost / 100 runs | $0.0046884 | $0.0082998 | +77.0% |",
        "",
        "## 10. Release decision",
        "",
        "![Release decision](../assets/diagrams/04_release_decision.png)",
        "",
        "The candidate remains **Review**. One regression divided by the 91 cases V1 classified correctly equals 1.10%, which is above the preregistered 1% pass bar. Aggregate improvement does not override the gate.",
        "",
        "## 11. Separate workbook-aligned result",
        "",
        "![Workbook-aligned comparison](../assets/figures/07_official_controlled_comparison.png)",
        "",
        "Controlled V1 scored 145/150 and focused V2 scored 146/150, with one win, zero regressions, and four remaining failures. The candidate remains Review because `other` precision is 0.882, below the 0.90 class-level floor.",
        "",
        "## 12. What is established",
        "",
        "- The supplied workbook contains four coherent policy-boundary failure mechanisms.",
        "- Notebook V2 materially improves its frozen development set, with complete paired accounting.",
        "- The regression and remaining notebook failure share a product-description boundary.",
        "- The workbook-aligned V2 produces a modest +1 result with no regressions.",
        "",
        "The evidence does **not** establish a causal 121-to-146 improvement, hidden-test performance, production readiness, or a calibrated LLM judge. Later benchmark-informed and incomplete runs are excluded.",
        "",
        "## 13. Final recommendation",
        "",
        "Resolve the routing contract, freeze the candidate and model configuration, commission an independently authored and adjudicated holdout, run once without further tuning, and require the regression, class-floor, latency, and cost gates to pass before release. After that, shadow deploy with human overrides, transfer rate, time-to-correct-queue, resolution, latency, and cost monitoring.",
        "",
        "## 14. Evidence index",
        "",
        "- [Completed workbook](../workbook/Week4_Customer_Support_Evaluation_Kajal_completed.xlsx)",
        "- [Executed notebook](../notebook/week4_customer_support_evals_executed.ipynb)",
        "- [Paired V1/V2 outcomes](../evidence/results/baseline_vs_improved.csv)",
        "- [Validation summary](../evidence/results/validation_summary.json)",
        "- [Workbook-aligned comparison](../evidence/results/official150_case_comparison.csv)",
        "- [Prompt versions](../prompts/)",
        "- [Integrity metadata](../metadata/)",
        f"- [LangSmith experiment comparison]({LANGSMITH_URL})",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main() -> None:
    metrics = get_metrics()
    wb_rows, wb_metrics, wb_matrix = workbook_baseline()
    ensure_assets(metrics, wb_metrics, wb_matrix)
    docx = build_docx(metrics, wb_rows, wb_metrics)
    markdown = build_markdown(metrics)
    print(f"DOCX={docx}")
    print(f"MARKDOWN={markdown}")
    print(f"FIGURES={len(list(FIGURES.glob('*.png')))}")


if __name__ == "__main__":
    main()
